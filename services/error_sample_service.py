"""
错误样本服务模块 (US-19)

提供错误样本的收集、查询、分类和管理功能。
支持从批量任务中自动收集错误，按类型/状态筛选，批量操作等。
"""
import uuid
import json
import os
from datetime import datetime
from typing import Optional, List, Dict, Any

from .database_service import AppDatabaseService
from .storage_service import StorageService


# 错误类型严重程度映射
ERROR_SEVERITY = {
    '识别错误-判断错误': 'high',
    '识别正确-判断错误': 'high',
    '缺失题目': 'high',
    'AI识别幻觉': 'high',
    '识别错误-判断正确': 'medium',
    '识别差异-判断正确': 'low',
    '识别题干-判断正确': 'low',
    '答案不匹配': 'medium'
}


class ErrorSampleService:
    """错误样本服务类"""
    
    # 内存缓存
    _cache: Dict[str, Any] = {}
    _cache_ttl = 300  # 5分钟
    
    @staticmethod
    def collect_from_task(task_id: str) -> Dict[str, Any]:
        """
        从批量任务中收集错误样本 (US-19.1)
        
        扫描任务中所有作业的 evaluation.errors，
        将错误信息存入 error_samples 表。
        
        Args:
            task_id: 批量任务ID
            
        Returns:
            dict: {collected: int, skipped: int, errors: list}
        """
        result = {'collected': 0, 'skipped': 0, 'errors': []}
        
        # 加载任务数据
        task_file = os.path.join(StorageService.BATCH_TASKS_DIR, f'{task_id}.json')
        if not os.path.exists(task_file):
            raise ValueError(f'任务不存在: {task_id}')

        try:
            with open(task_file, 'r', encoding='utf-8') as f:
                task_data = json.load(f)
        except Exception as e:
            raise ValueError(f'读取任务文件失败: {e}')
        
        # 遍历作业项
        for hw_item in task_data.get('homework_items', []):
            evaluation = hw_item.get('evaluation') or {}
            errors = evaluation.get('errors') or []
            
            if not errors:
                continue
            
            homework_id = hw_item.get('homework_id', '')
            book_id = hw_item.get('book_id', '')
            book_name = hw_item.get('book_name', '')
            page_num = hw_item.get('page_num')
            pic_path = hw_item.get('pic_path', '')
            matched_dataset = hw_item.get('matched_dataset', '')
            
            # 推断学科ID
            subject_id = ErrorSampleService._infer_subject_id(book_name)
            
            for error in errors:
                sample_id = str(uuid.uuid4())[:8]
                question_index = error.get('index', '')
                error_type = error.get('error_type', '其他')
                
                # 检查是否已存在相同样本
                existing = AppDatabaseService.execute_one(
                    """SELECT id FROM error_samples 
                       WHERE task_id = %s AND homework_id = %s AND question_index = %s""",
                    (task_id, homework_id, question_index)
                )
                
                if existing:
                    result['skipped'] += 1
                    continue
                
                # 插入样本
                try:
                    sql = """
                        INSERT INTO error_samples 
                        (sample_id, task_id, homework_id, dataset_id, book_id, book_name,
                         page_num, question_index, subject_id, error_type, base_answer,
                         base_user, hw_user, pic_path, status)
                        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                    """
                    AppDatabaseService.execute_insert(sql, (
                        sample_id, task_id, homework_id, matched_dataset,
                        book_id, book_name, page_num, question_index, subject_id,
                        error_type, error.get('base_answer', ''),
                        error.get('base_user', ''), error.get('hw_user', ''),
                        pic_path, 'pending'
                    ))
                    result['collected'] += 1
                except Exception as e:
                    result['errors'].append(f'插入样本失败: {e}')
        
        # 清除缓存
        ErrorSampleService._cache.clear()
        
        return result

    
    @staticmethod
    def get_samples(
        page: int = 1,
        page_size: int = 20,
        error_type: str = None,
        status: str = None,
        subject_id: int = None,
        cluster_id: str = None,
        task_id: str = None,
        keyword: str = None
    ) -> Dict[str, Any]:
        """
        获取错误样本列表 (US-19.2)
        
        支持多条件筛选和分页。
        
        Args:
            page: 页码
            page_size: 每页数量
            error_type: 错误类型筛选
            status: 状态筛选
            subject_id: 学科筛选
            cluster_id: 聚类筛选
            task_id: 任务筛选
            keyword: 关键词搜索
            
        Returns:
            dict: {items: list, total: int, page: int, page_size: int}
        """
        # 构建查询条件
        where_clauses = ['1=1']
        params = []
        
        if error_type:
            where_clauses.append('error_type = %s')
            params.append(error_type)
        
        if status:
            where_clauses.append('status = %s')
            params.append(status)
        
        if subject_id is not None:
            where_clauses.append('subject_id = %s')
            params.append(subject_id)
        
        if cluster_id:
            where_clauses.append('cluster_id = %s')
            params.append(cluster_id)
        
        if task_id:
            where_clauses.append('task_id = %s')
            params.append(task_id)
        
        if keyword:
            where_clauses.append(
                '(book_name LIKE %s OR question_index LIKE %s OR base_answer LIKE %s)'
            )
            kw = f'%{keyword}%'
            params.extend([kw, kw, kw])
        
        where_sql = ' AND '.join(where_clauses)
        
        # 查询总数
        count_sql = f'SELECT COUNT(*) as total FROM error_samples WHERE {where_sql}'
        count_result = AppDatabaseService.execute_one(count_sql, tuple(params) if params else None)
        total = count_result['total'] if count_result else 0
        
        # 查询列表
        offset = (page - 1) * page_size
        list_sql = f"""
            SELECT sample_id, task_id, homework_id, dataset_id, book_id, book_name,
                   page_num, question_index, subject_id, error_type, base_answer,
                   base_user, hw_user, pic_path, status, notes, cluster_id,
                   created_at, updated_at
            FROM error_samples 
            WHERE {where_sql}
            ORDER BY created_at DESC
            LIMIT %s OFFSET %s
        """
        params.extend([page_size, offset])
        rows = AppDatabaseService.execute_query(list_sql, tuple(params))
        
        # 格式化结果
        items = []
        for row in rows:
            items.append({
                'sample_id': row['sample_id'],
                'task_id': row['task_id'],
                'homework_id': row['homework_id'],
                'dataset_id': row.get('dataset_id'),
                'book_id': row.get('book_id'),
                'book_name': row.get('book_name', ''),
                'page_num': row.get('page_num'),
                'question_index': row['question_index'],
                'subject_id': row.get('subject_id'),
                'error_type': row['error_type'],
                'severity': ERROR_SEVERITY.get(row['error_type'], 'medium'),
                'base_answer': row.get('base_answer', ''),
                'base_user': row.get('base_user', ''),
                'hw_user': row.get('hw_user', ''),
                'pic_path': row.get('pic_path', ''),
                'status': row['status'],
                'notes': row.get('notes', ''),
                'cluster_id': row.get('cluster_id'),
                'created_at': row['created_at'].isoformat() if row.get('created_at') else '',
                'updated_at': row['updated_at'].isoformat() if row.get('updated_at') else ''
            })
        
        return {
            'items': items,
            'total': total,
            'page': page,
            'page_size': page_size,
            'total_pages': (total + page_size - 1) // page_size
        }

    
    @staticmethod
    def get_sample_detail(sample_id: str) -> Optional[Dict[str, Any]]:
        """
        获取样本详情 (US-19.3)
        
        Args:
            sample_id: 样本ID
            
        Returns:
            dict: 样本详细信息
        """
        sql = """
            SELECT sample_id, task_id, homework_id, dataset_id, book_id, book_name,
                   page_num, question_index, subject_id, error_type, base_answer,
                   base_user, hw_user, pic_path, status, notes, cluster_id,
                   created_at, updated_at
            FROM error_samples 
            WHERE sample_id = %s
        """
        row = AppDatabaseService.execute_one(sql, (sample_id,))
        
        if not row:
            return None
        
        return {
            'sample_id': row['sample_id'],
            'task_id': row['task_id'],
            'homework_id': row['homework_id'],
            'dataset_id': row.get('dataset_id'),
            'book_id': row.get('book_id'),
            'book_name': row.get('book_name', ''),
            'page_num': row.get('page_num'),
            'question_index': row['question_index'],
            'subject_id': row.get('subject_id'),
            'error_type': row['error_type'],
            'severity': ERROR_SEVERITY.get(row['error_type'], 'medium'),
            'base_answer': row.get('base_answer', ''),
            'base_user': row.get('base_user', ''),
            'hw_user': row.get('hw_user', ''),
            'pic_path': row.get('pic_path', ''),
            'status': row['status'],
            'notes': row.get('notes', ''),
            'cluster_id': row.get('cluster_id'),
            'created_at': row['created_at'].isoformat() if row.get('created_at') else '',
            'updated_at': row['updated_at'].isoformat() if row.get('updated_at') else ''
        }
    
    @staticmethod
    def update_status(
        sample_ids: List[str],
        status: str,
        notes: str = None
    ) -> int:
        """
        批量更新样本状态 (US-19.4)
        
        Args:
            sample_ids: 样本ID列表
            status: 新状态
            notes: 备注（可选）
            
        Returns:
            int: 更新的记录数
        """
        if not sample_ids:
            return 0
        
        valid_statuses = ['pending', 'analyzed', 'fixed', 'ignored']
        if status not in valid_statuses:
            raise ValueError(f'无效状态: {status}')
        
        placeholders = ','.join(['%s'] * len(sample_ids))
        
        if notes:
            sql = f"""
                UPDATE error_samples 
                SET status = %s, notes = %s, updated_at = NOW()
                WHERE sample_id IN ({placeholders})
            """
            params = [status, notes] + sample_ids
        else:
            sql = f"""
                UPDATE error_samples 
                SET status = %s, updated_at = NOW()
                WHERE sample_id IN ({placeholders})
            """
            params = [status] + sample_ids
        
        result = AppDatabaseService.execute_update(sql, tuple(params))
        
        # 清除缓存
        ErrorSampleService._cache.clear()
        
        return result
    
    @staticmethod
    def get_statistics() -> Dict[str, Any]:
        """
        获取错误样本统计
        
        Returns:
            dict: {total, by_status, by_error_type, by_subject}
        """
        # 总数和状态分布
        status_sql = """
            SELECT status, COUNT(*) as count 
            FROM error_samples 
            GROUP BY status
        """
        status_rows = AppDatabaseService.execute_query(status_sql)
        by_status = {row['status']: row['count'] for row in status_rows}
        total = sum(by_status.values())
        
        # 错误类型分布
        type_sql = """
            SELECT error_type, COUNT(*) as count 
            FROM error_samples 
            GROUP BY error_type
            ORDER BY count DESC
        """
        type_rows = AppDatabaseService.execute_query(type_sql)
        by_error_type = {row['error_type']: row['count'] for row in type_rows}
        
        # 学科分布
        subject_sql = """
            SELECT subject_id, COUNT(*) as count 
            FROM error_samples 
            WHERE subject_id IS NOT NULL
            GROUP BY subject_id
        """
        subject_rows = AppDatabaseService.execute_query(subject_sql)
        by_subject = {row['subject_id']: row['count'] for row in subject_rows}
        
        return {
            'total': total,
            'by_status': by_status,
            'by_error_type': by_error_type,
            'by_subject': by_subject
        }

    
    @staticmethod
    def export_samples(
        filters: Dict[str, Any] = None,
        format: str = 'xlsx'
    ) -> str:
        """
        导出错误样本 (US-22.5)
        
        Args:
            filters: 筛选条件
            format: 导出格式 xlsx/csv
            
        Returns:
            str: 导出文件路径
        """
        from openpyxl import Workbook
        import csv
        
        filters = filters or {}
        
        # 获取所有符合条件的样本
        all_samples = []
        page = 1
        while True:
            result = ErrorSampleService.get_samples(
                page=page,
                page_size=1000,
                error_type=filters.get('error_type'),
                status=filters.get('status'),
                subject_id=filters.get('subject_id'),
                task_id=filters.get('task_id')
            )
            all_samples.extend(result['items'])
            if page >= result['total_pages']:
                break
            page += 1
        
        # 确保导出目录存在
        export_dir = StorageService.EXPORTS_DIR
        StorageService.ensure_dir(export_dir)
        
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        
        if format == 'csv':
            filename = f'error_samples_{timestamp}.csv'
            filepath = os.path.join(export_dir, filename)
            
            with open(filepath, 'w', newline='', encoding='utf-8-sig') as f:
                writer = csv.writer(f)
                # 写入表头
                writer.writerow([
                    '样本ID', '任务ID', '作业ID', '书本名称', '页码', '题号',
                    '错误类型', '严重程度', '基准答案', '基准用户答案', 'AI答案',
                    '状态', '备注', '创建时间'
                ])
                # 写入数据
                for sample in all_samples:
                    writer.writerow([
                        sample['sample_id'], sample['task_id'], sample['homework_id'],
                        sample['book_name'], sample['page_num'], sample['question_index'],
                        sample['error_type'], sample['severity'],
                        sample['base_answer'], sample['base_user'], sample['hw_user'],
                        sample['status'], sample['notes'], sample['created_at']
                    ])
        else:
            filename = f'error_samples_{timestamp}.xlsx'
            filepath = os.path.join(export_dir, filename)
            
            wb = Workbook()
            ws = wb.active
            ws.title = '错误样本'
            
            # 写入表头
            headers = [
                '样本ID', '任务ID', '作业ID', '书本名称', '页码', '题号',
                '错误类型', '严重程度', '基准答案', '基准用户答案', 'AI答案',
                '状态', '备注', '创建时间'
            ]
            ws.append(headers)
            
            # 写入数据
            for sample in all_samples:
                ws.append([
                    sample['sample_id'], sample['task_id'], sample['homework_id'],
                    sample['book_name'], sample['page_num'], sample['question_index'],
                    sample['error_type'], sample['severity'],
                    sample['base_answer'], sample['base_user'], sample['hw_user'],
                    sample['status'], sample['notes'], sample['created_at']
                ])
            
            wb.save(filepath)
        
        return filepath
    
    @staticmethod
    def _infer_subject_id(book_name: str) -> Optional[int]:
        """从书名推断学科ID"""
        if not book_name:
            return None
        
        book_name_lower = book_name.lower()
        subject_keywords = {
            0: ['英语', 'english'],
            1: ['语文', 'chinese'],
            2: ['数学', 'math'],
            3: ['物理', 'physics'],
            4: ['化学', 'chemistry'],
            5: ['生物', 'biology'],
            6: ['地理', 'geography']
        }
        
        for subject_id, keywords in subject_keywords.items():
            for keyword in keywords:
                if keyword in book_name_lower:
                    return subject_id
        
        return None
